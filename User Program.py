"""
laptop_gui_server.py
=====================================================
Laptop GUI สำหรับระบบ PLC Vision Inspection

สถาปัตยกรรม:
- Raspberry Pi = master (เก็บ SQLite + รูปภาพทั้งหมด, เปิด Read API พอร์ต 5001)
- Laptop (ไฟล์นี้) = GUI ล้วนๆ ไม่มี SQLite/ไฟล์รูปของตัวเอง
    1) รับ push ข้อมูล real-time จาก Pi (POST /api/plc_data) -> เก็บใน memory cache
    2) ดูประวัติ/ค้นหาย้อนหลัง -> proxy ไป Pi ผ่าน /api/history
    3) ดูรูปภาพ -> proxy ไป Pi ผ่าน /api/image/<item_number>
    4) Export Excel -> ดึงข้อมูลทั้งหมดจาก Pi มาสร้างไฟล์ .xlsx ให้ดาวน์โหลด

ก่อนรัน ต้องติดตั้ง: pip install flask requests openpyxl
=====================================================
"""

import io
import threading
from collections import deque, Counter
from datetime import datetime

import requests
from flask import Flask, request, jsonify, Response, send_file, render_template_string
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIGURATION
# ============================================================
# !!! ระบุ IP จริงของ Raspberry Pi ที่นี่ !!!
PI_IP = "192.168.3.100" 
PI_API_PORT = 5001
PI_API_BASE = f"http://{PI_IP}:{PI_API_PORT}"

LAPTOP_LISTEN_HOST = "0.0.0.0"
LAPTOP_LISTEN_PORT = 5000

LIVE_CACHE_MAXLEN = 200
PI_REQUEST_TIMEOUT = 5  # seconds

app = Flask(__name__)

# ============================================================
# IN-MEMORY LIVE CACHE (ไม่มี SQLite ฝั่ง Laptop)
# ============================================================
live_cache = deque(maxlen=LIVE_CACHE_MAXLEN)
live_cache_lock = threading.Lock()


# ============================================================
# 1) รับ PUSH ข้อมูล real-time จาก Pi
# ============================================================
@app.route('/api/plc_data', methods=['POST'])
def receive_data():
    data = request.get_json()
    if not data:
        return jsonify({"status": "failed"}), 400

    record = {
        "timestamp": data.get('timestamp'),
        "item_number": data.get('item_number'),
        "shape_code": data.get('shape_code'),
        "shape_name": data.get('shape_name'),
        "error_val": float(data.get('error_val', 0.0)),
        "rpm": int(data.get('rpm', 0)),
        "set_hz": int(data.get('set_hz', 0)),
        "speed_ms": float(data.get('speed_ms', 0.0)),
        "cycle_speed": float(data.get('cycle_speed', 0.0)),
        "image_path": data.get('image_path', ''),
    }

    with live_cache_lock:
        live_cache.appendleft(record)

    print(f"[LIVE] Item #{record['item_number']} | {record['shape_name']} | Img: {record['image_path'] or '-'}")
    return jsonify({"status": "success"}), 200


@app.route('/api/live')
def api_live():
    with live_cache_lock:
        data = list(live_cache)
    
    # คำนวณสรุปจำนวนนับแยกตาม Class จาก cache ที่มีอยู่
    class_counts = Counter(item['shape_name'] for item in data)
    
    return jsonify({
        "status": "success", 
        "count": len(data), 
        "class_counts": dict(class_counts),
        "data": data
    })


# ============================================================
# 2) PROXY ประวัติย้อนหลังไปหา Pi
# ============================================================
@app.route('/api/history')
def api_history_proxy():
    limit = request.args.get('limit', default=100, type=int)
    offset = request.args.get('offset', default=0, type=int)
    try:
        resp = requests.get(
            f"{PI_API_BASE}/api/history",
            params={"limit": limit, "offset": offset},
            timeout=PI_REQUEST_TIMEOUT
        )
        resp.raise_for_status()
        return jsonify(resp.json())
    except requests.exceptions.RequestException as e:
        return jsonify({"status": "error", "message": f"ติดต่อ Pi ที่ {PI_API_BASE} ไม่ได้: {e}"}), 502


# ============================================================
# 3) PROXY รูปภาพไปหา Pi
# ============================================================
@app.route('/api/image/<int:item_number>')
def api_image_proxy(item_number):
    try:
        resp = requests.get(
            f"{PI_API_BASE}/api/image/{item_number}",
            timeout=PI_REQUEST_TIMEOUT
        )
        if resp.status_code != 200:
            return jsonify({"status": "error", "message": "ไม่พบรูปภาพสำหรับ item นี้บน Pi"}), 404
        return Response(resp.content, mimetype='image/jpeg')
    except requests.exceptions.RequestException as e:
        return jsonify({"status": "error", "message": f"ติดต่อ Pi ไม่ได้: {e}"}), 502


# ============================================================
# 4) EXPORT EXCEL (ดึงข้อมูลทั้งหมดจาก Pi มาสร้างไฟล์ .xlsx)
# ============================================================
@app.route('/api/export')
def api_export_excel():
    all_rows = []
    offset = 0
    page_size = 1000
    try:
        while True:
            resp = requests.get(
                f"{PI_API_BASE}/api/history",
                params={"limit": page_size, "offset": offset},
                timeout=PI_REQUEST_TIMEOUT
            )
            resp.raise_for_status()
            payload = resp.json()
            rows = payload.get("data", [])
            if not rows:
                break
            all_rows.extend(rows)
            offset += page_size
            if len(rows) < page_size:
                break
    except requests.exceptions.RequestException as e:
        return jsonify({"status": "error", "message": f"ดึงข้อมูลจาก Pi เพื่อ Export ไม่สำเร็จ: {e}"}), 502

    wb = Workbook()
    ws = wb.active
    ws.title = "PLC Vision Logs"

    headers = [
        "ID", "Timestamp", "Item Number", "Shape Code", 
        "Shape Name", "Error (%)", "RPM", "Set Hz", 
        "Synchronous Speed", "Slip %", "Speed (m/s)", "Cycle Speed", "Image File"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for r in all_rows:
        set_hz = r.get("set_hz") or 0
        rpm = r.get("rpm") or 0
        
        # คำนวณ Synchronous Speed = (120 * Set Hz) / 4
        sync_speed = (120.0 * set_hz) / 4.0 if set_hz else 0.0
        
        # คำนวณ Slip % = ((Sync Speed - RPM) / Sync Speed) * 100
        if sync_speed > 0:
            slip_percent = ((sync_speed - rpm) / sync_speed) * 100.0
        else:
            slip_percent = 0.0

        ws.append([
            r.get("id"),
            r.get("timestamp"),
            r.get("item_number"),
            r.get("shape_code"),
            r.get("shape_name"),
            r.get("error_val"),
            rpm,
            set_hz,
            round(sync_speed, 2),
            round(slip_percent, 2),
            r.get("speed_ms"),
            r.get("cycle_speed"),
            r.get("image_path")
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"plc_vision_logs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


# ============================================================
# HTML GUI TEMPLATE (Dark Modern UI พร้อมส่วนแสดง Class Counters)
# ============================================================
HTML_PAGE = """
<!DOCTYPE html>
<html lang="th">
<head>
<meta charset="UTF-8">
<title>PLC Vision Inspection - Laptop GUI</title>
<style>
  :root {
    --bg: #0f172a;
    --card-bg: #1e293b;
    --border: #334155;
    --text: #f8fafc;
    --muted: #94a3b8;
    --primary: #3b82f6;
    --primary-hover: #2563eb;
    --good: #22c55e;
    --bad: #ef4444;
    --warning: #f59e0b;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; }
  body { background-color: var(--bg); color: var(--text); padding: 20px; }
  header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; border-bottom: 1px solid var(--border); padding-bottom: 15px; }
  h1 { font-size: 1.5rem; color: var(--text); display: flex; align-items: center; gap: 10px; }
  .badge { background: var(--primary); color: white; padding: 4px 10px; border-radius: 12px; font-size: 0.8rem; }
  .btn { background: var(--primary); color: white; border: none; padding: 8px 16px; border-radius: 6px; cursor: pointer; font-weight: 600; transition: background 0.2s; }
  .btn:hover { background: var(--primary-hover); }
  .btn-success { background: var(--good); }
  .btn-success:hover { background: #16a34a; }
  
  /* Counter Summary Panel */
  .counter-container { display: grid; grid-template-columns: repeat(auto-fit, minmax(140px, 1fr)); gap: 12px; margin-bottom: 20px; }
  .counter-card { background: var(--card-bg); border: 1px solid var(--border); border-radius: 8px; padding: 12px; text-align: center; }
  .counter-title { font-size: 0.85rem; color: var(--muted); margin-bottom: 6px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
  .counter-value { font-size: 1.4rem; font-weight: bold; color: var(--text); }

  .grid-2 { display: grid; grid-template-columns: 1fr; gap: 20px; }
  @media(min-width: 1024px) { .grid-2 { grid-template-columns: 1fr 1fr; } }
  .card { background: var(--card-bg); border: 1px solid var(--border); border-radius: 10px; padding: 20px; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.1); }
  .card h2 { font-size: 1.1rem; margin-bottom: 15px; display: flex; justify-content: space-between; align-items: center; }
  
  table { width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 0.9rem; }
  th, td { padding: 10px 12px; text-align: left; border-bottom: 1px solid var(--border); }
  th { background-color: rgba(51, 65, 85, 0.4); color: var(--muted); font-weight: 600; }
  tr:hover { background-color: rgba(255, 255, 255, 0.02); }
  .tag { padding: 3px 8px; border-radius: 4px; font-size: 0.75rem; font-weight: bold; }
  .tag-good { background: rgba(34, 197, 94, 0.2); color: var(--good); }
  .tag-bad { background: rgba(239, 68, 68, 0.2); color: var(--bad); }
  
  .search-box { display: flex; gap: 10px; margin-bottom: 15px; }
  .search-box input { flex: 1; background: var(--bg); border: 1px solid var(--border); border-radius: 6px; padding: 8px 12px; color: white; }
  
  .modal-bg { display: none; position: fixed; top: 0; left: 0; width: 100%; height: 100%; background: rgba(0,0,0,0.7); justify-content: center; align-items: center; z-index: 1000; }
  .modal-content { background: var(--card-bg); padding: 20px; border-radius: 10px; max-width: 600px; width: 90%; border: 1px solid var(--border); position: relative; text-align: center; }
  .modal-content img { max-width: 100%; height: auto; border-radius: 6px; margin-top: 15px; border: 1px solid var(--border); }
  .close-btn { position: absolute; top: 15px; right: 15px; background: none; border: none; color: var(--muted); font-size: 1.2rem; cursor: pointer; }
</style>
</head>
<body>

<header>
  <h1>PLC Vision Inspection <span class="badge">Laptop GUI (Live)</span></h1>
  <div>
    <button class="btn btn-success" onclick="exportExcel()">📥 Export Excel (Pi)</button>
  </div>
</header>

<!-- CLASS COUNTERS REAL-TIME PANEL -->
<div id="counter-panel" class="counter-container">
  <div class="counter-card">
    <div class="counter-title">กำลังโหลด...</div>
    <div class="counter-value">-</div>
  </div>
</div>

<div class="grid-2">
  <!-- LIVE FEED CARD -->
  <div class="card">
    <h2><span>🔴 Real-Time Feed (จาก Pi)</span> <span id="live-count" class="badge">0 รายการ</span></h2>
    <div style="overflow-x: auto; max-height: 450px; overflow-y: auto;">
      <table>
        <thead>
          <tr>
            <th>เวลา</th>
            <th>Item #</th>
            <th>ชื่อชิ้นงาน (Class)</th>
            <th>Error</th>
            <th>รูปภาพ</th>
          </tr>
        </thead>
        <tbody id="live-tbody">
          <tr><td colspan="5" style="text-align:center; color:var(--muted);">รอข้อมูล Real-Time จาก Pi...</td></tr>
        </tbody>
      </table>
    </div>
  </div>

  <!-- HISTORY SEARCH CARD -->
  <div class="card">
    <h2><span>🔍 ค้นหาประวัติย้อนหลัง (จาก Pi SQLite)</span></h2>
    <div class="search-box">
      <input type="number" id="search-limit" placeholder="จำนวนแถว (Default 50)" value="50">
      <button class="btn" onclick="loadHistory()">ค้นหา</button>
    </div>
    <div style="overflow-x: auto; max-height: 450px; overflow-y: auto;">
      <table>
        <thead>
          <tr>
            <th>ID</th>
            <th>เวลา</th>
            <th>Item #</th>
            <th>ชื่อชิ้นงาน (Class)</th>
            <th>Error</th>
            <th>ความเร็ว</th>
            <th>รูป</th>
          </tr>
        </thead>
        <tbody id="history-tbody">
          <tr><td colspan="7" style="text-align:center; color:var(--muted);">กดปุ่มค้นหาเพื่อดูประวัติ</td></tr>
        </tbody>
      </table>
    </div>
  </div>
</div>

<!-- IMAGE MODAL -->
<div id="modal-bg" class="modal-bg" onclick="closeModalIfBg(event)">
  <div class="modal-content">
    <button class="close-btn" onclick="closeModal()">✖</button>
    <h3 id="modal-title">Item Image</h3>
    <img id="modal-img" src="" alt="Item Image">
  </div>
</div>

<script>
  async function refreshLive() {
    try {
      const res = await fetch('/api/live');
      const json = await res.json();
      if (json.status === 'success') {
        document.getElementById('live-count').textContent = `${json.count} รายการ`;
        
        // อัปเดตแผงนับจำนวน Class แบบ Real-Time
        renderCounters(json.class_counts);

        const tbody = document.getElementById('live-tbody');
        if (json.data.length > 0) {
          tbody.innerHTML = json.data.map(r => `
            <tr>
              <td>${r.timestamp || '-'}</td>
              <td><b>#${r.item_number}</b></td>
              <td>${r.shape_name}</td>
              <td><span class="tag ${r.error_val > 5 ? 'tag-bad' : 'tag-good'}">${r.error_val.toFixed(1)}%</span></td>
              <td>${r.image_path ? `<button class="btn" style="padding:2px 8px; font-size:0.8rem;" onclick="showImage(${r.item_number})">ดูรูป</button>` : '-'}</td>
            </tr>
          `).join('');
        } else {
          tbody.innerHTML = '<tr><td colspan="5" style="text-align:center; color:var(--muted);">ยังไม่มีข้อมูลส่งเข้ามา</td></tr>';
        }
      }
    } catch (e) {
      console.error("Fetch live error:", e);
    }
  }

  function renderCounters(classCounts) {
    const container = document.getElementById('counter-panel');
    if (!classCounts || Object.keys(classCounts).length === 0) {
      container.innerHTML = `<div class="counter-card"><div class="counter-title">ยังไม่มีข้อมูลนับ</div><div class="counter-value">0</div></div>`;
      return;
    }
    
    let html = '';
    for (const [className, count] of Object.entries(classCounts)) {
      html += `
        <div class="counter-card">
          <div class="counter-title" title="${className}">${className}</div>
          <div class="counter-value">${count}</div>
        </div>
      `;
    }
    container.innerHTML = html;
  }

  async function loadHistory() {
    const limit = document.getElementById('search-limit').value || 50;
    const tbody = document.getElementById('history-tbody');
    tbody.innerHTML = '<tr><td colspan="7" style="text-align:center; color:var(--muted);">กำลังโหลด...</td></tr>';
    try {
      const res = await fetch(`/api/history?limit=${limit}`);
      const json = await res.json();
      if (json.status === 'success' && json.data.length > 0) {
        tbody.innerHTML = json.data.map(r => `
          <tr>
            <td>${r.id}</td>
            <td>${r.timestamp}</td>
            <td><b>#${r.item_number}</b></td>
            <td>${r.shape_name}</td>
            <td><span class="tag ${r.error_val > 5 ? 'tag-bad' : 'tag-good'}">${r.error_val.toFixed(1)}%</span></td>
            <td>${r.speed_ms} m/s</td>
            <td>${r.image_path ? `<button class="btn" style="padding:2px 8px; font-size:0.8rem;" onclick="showImage(${r.item_number})">ดูรูป</button>` : '-'}</td>
          </tr>
        `).join('');
      } else {
        tbody.innerHTML = '<tr><td colspan="7" style="text-align:center; color:var(--muted);">ไม่พบประวัติข้อมูลใน Pi</td></tr>';
      }
    } catch (e) {
      tbody.innerHTML = `<tr><td colspan="7" style="color:var(--bad); text-align:center;">โหลดข้อมูลจาก Pi ไม่สำเร็จ: ${e.message}</td></tr>`;
    }
  }

  function showImage(itemNumber) {
    if(!itemNumber) return;
    document.getElementById('modal-title').textContent = `Item #${itemNumber}`;
    document.getElementById('modal-img').src = `/api/image/${itemNumber}?t=${Date.now()}`;
    document.getElementById('modal-bg').style.display = 'flex';
  }
  function closeModal() { document.getElementById('modal-bg').style.display = 'none'; }
  function closeModalIfBg(e) { if (e.target.id === 'modal-bg') closeModal(); }

  function exportExcel() {
    window.location.href = '/api/export';
  }

  refreshLive();
  loadHistory();
  setInterval(refreshLive, 1500);
</script>

</body>
</html>
"""


@app.route('/')
def index():
    return render_template_string(HTML_PAGE)


# ============================================================
# MAIN
# ============================================================
if __name__ == '__main__':
    print(f"[GUI] Laptop GUI Server starting on {LAPTOP_LISTEN_HOST}:{LAPTOP_LISTEN_PORT} ...")
    app.run(host=LAPTOP_LISTEN_HOST, port=LAPTOP_LISTEN_PORT, debug=False)